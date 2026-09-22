import io
import json
import re
import sqlite3
import tempfile
import unittest
from html import unescape
from datetime import date
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from src import web_ui as ui, team_directory as d
from src.app_request_state import ensure_app_tables
from src.db_schema import ensure_core_tables

ENV = {'WORK_ON_HOLIDAY_SUPERUSER_LOGIN': 'root', 'WORK_ON_HOLIDAY_SUPERUSER_PASSWORD': 'test-password'}
ROWS = [['Руководитель Главный', 'root@example.org', ''],
        ['Начальник Средний', 'middle@example.org', 'root@example.org'],
        ['Сотрудник Нижний', 'worker@example.org', 'middle@example.org'],
        ['Сотрудник Прямой', 'direct@example.org', 'root@example.org'],
        ['Чужой Сотрудник', 'other@example.org', '']]


class TeamTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / 'db.sqlite'
        self.path_patch = patch.object(ui, 'DB_PATH', self.path)
        self.path_patch.start()
        self.env_patch = patch.dict('os.environ', ENV)
        self.env_patch.start()
        ui.ensure_app_tables_for_app()
        self.client = TestClient(ui.app)
        with ui.get_db_connection() as conn:
            d.apply_import(conn, d.plan_import(conn, ROWS))
            self.keys = {p['email']: p['key'] for p in d.employees(conn).values()}
            for mail, key in self.keys.items():
                ui.upsert_employee_token(conn, key, mail)
        self.root = self.keys['root@example.org']
        self.worker = self.keys['worker@example.org']

    def tearDown(self):
        self.client.close()
        self.path_patch.stop()
        self.env_patch.stop()
        self.tmp.cleanup()

    def login(self, mail='root@example.org'):
        self.client.cookies.set(ui.EMPLOYEE_TOKEN_COOKIE_NAME, mail)

    def admin(self):
        self.client.post('/superuser/login', data={'superuser_login':'root','superuser_password':'test-password'})

    def task(self, key, identifier, day='2026-09-12', text='Уникальная задача'):
        with ui.get_db_connection() as conn:
            conn.execute('''INSERT INTO survey_responses(response_id,full_name_key,full_name,request_type,
                planned_work_date,planned_work_time,task_description) VALUES (?,?,?,'Подать заявку',?,'10:00 - 12:00',?)''',
                (identifier, key, ui.get_employee_display_name(conn, key), day, text))

    def exported_rows(self, response):
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
        self.assertIn('spreadsheetml.sheet', response.headers['content-type'])
        self.assertEqual(response.headers['cache-control'], 'private, no-store')
        self.assertIn('attachment;', response.headers['content-disposition'])
        book = load_workbook(io.BytesIO(response.content))
        self.addCleanup(book.close)
        values = list(book['Заявки'].values)
        return book, [dict(zip(values[0], row)) for row in values[1:]]

    def export_link(self, page):
        self.assertEqual(page.status_code, 200)
        return unescape(re.search(r'id="request-export" href="([^"]+)"', page.text)[1])

    def test_superuser_heads_import_and_legacy_structure_without_fake_employee(self):
        with ui.get_db_connection() as conn:
            legacy = ui.register_employee_directory_entry(conn, 'Старый Сотрудник')['employee_key']
            ui.upsert_employee_token(conn, legacy, 'legacy-token')
            # A database from before this change has NULL managers at the top.
            conn.execute('UPDATE app_org_employee SET manager_key=NULL WHERE employee_key=?', (self.root,))
            before = ui.get_employee_token_record(conn, legacy)['token_hash']
            d.ensure_team_tables(conn)
            people = d.employees(conn)
            self.assertEqual(people[legacy]['manager'], d.SUPERUSER_KEY)
            self.assertEqual(people[self.root]['manager'], d.SUPERUSER_KEY)
            self.assertEqual(people[self.worker]['manager'], self.keys['middle@example.org'])
            self.assertEqual(d.team_keys(conn, d.SUPERUSER_KEY),
                             {self.root, self.keys['other@example.org'], legacy})
            self.assertEqual(d.team_keys(conn, d.SUPERUSER_KEY, 'all'), set(people))
            self.assertNotIn(d.SUPERUSER_KEY, people)
            self.assertIsNone(conn.execute('SELECT 1 FROM app_employee_directory WHERE full_name_key=?',
                                          (d.SUPERUSER_KEY,)).fetchone())
            plan = d.plan_import(conn, [['Начальник Средний', 'middle@example.org', '']])
            self.assertFalse(plan['errors'])
            self.assertEqual(plan['changes'][0]['manager'], d.SUPERUSER_KEY)
            d.apply_import(conn, plan)
            self.assertIn(self.keys['middle@example.org'], d.team_keys(conn, d.SUPERUSER_KEY))
            self.assertNotIn(self.worker, d.team_keys(conn, self.root, 'all'))
            self.assertIn(self.worker, d.team_keys(conn, d.SUPERUSER_KEY, 'all'))
            self.assertEqual(ui.get_employee_token_record(conn, legacy)['token_hash'], before)
        overview = {p['employee_key']: p for p in ui.get_admin_employees_overview()}
        self.assertEqual(overview[legacy]['manager_name'], 'Суперпользователь')
        self.assertEqual(overview[self.worker]['manager_name'], 'Начальник Средний')

    def test_superuser_default_all_levels_details_and_filtered_export(self):
        self.task(self.root, 1, day=date.today().isoformat(), text='Задача верхнего уровня')
        self.task(self.worker, 2, day=date.today().isoformat(), text='Задача глубоко в структуре')
        self.task(self.keys['other@example.org'], 3, day=date.today().isoformat(), text='Другая ветвь')
        # Historical requests may be the only record for an employee.
        self.task('legacy-request-only', 4, day=date.today().isoformat(), text='Задача из старой базы')
        with ui.get_db_connection() as conn:
            conn.execute("UPDATE survey_responses SET full_name='Исторический Сотрудник' WHERE response_id=4")
        self.admin()
        page = self.client.get('/manager')
        self.assertEqual(page.status_code, 200)
        self.assertIn('value="all" selected', page.text)
        self.assertIn('Оргструктура · сотрудников: 6', page.text)
        _, rows = self.exported_rows(self.client.get(self.export_link(page)))
        self.assertEqual({r['ID заявки'] for r in rows}, {1, 2, 3, 4})
        self.assertEqual(next(r for r in rows if r['ID заявки'] == 4)['Руководитель'], 'Суперпользователь')
        for identifier in (1, 2, 3, 4):
            detail = self.client.get(f'/manager/requests/{identifier}')
            self.assertEqual(detail.status_code, 200)
            self.assertIn('Управлять заявками сотрудника', detail.text)
        self.assertEqual(self.client.get('/manager/requests/9999').status_code, 404)
        params = dict(apply='1', scope='direct', period='current')
        _, rows = self.exported_rows(self.client.get('/manager/requests/export', params=params))
        self.assertEqual({r['ID заявки'] for r in rows}, {1, 3, 4})
        params.update(scope='all', employees=[self.worker], surname='Нижний', text='глубоко')
        page = self.client.get('/manager', params=params)
        _, rows = self.exported_rows(self.client.get(self.export_link(page)))
        self.assertEqual([r['ID заявки'] for r in rows], [2])
        self.assertIn('Оргструктура · сотрудников: 6', page.text)
        self.exported_rows(self.client.get('/admin/requests/export'))

    def test_superuser_access_empty_organization_and_navigation(self):
        with ui.get_db_connection() as conn:
            conn.execute('DELETE FROM app_org_employee')
            conn.execute('DELETE FROM app_employee_directory')
            conn.execute('DELETE FROM app_employee_auth')
        self.admin()
        page = self.client.get('/manager')
        self.assertEqual(page.status_code, 200)
        self.assertIn('Оргструктура · сотрудников: 0', page.text)
        _, rows = self.exported_rows(self.client.get(self.export_link(page)))
        self.assertEqual(rows, [])
        for url in ('/', '/admin', '/admin/users', '/admin/requests', '/admin/test-data', '/employee?admin_mode=1'):
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200)
            self.assertIn('href="/manager"', response.text, url)

    def test_superuser_session_precedes_employee_and_filters_are_isolated(self):
        self.login('worker@example.org')
        self.task(self.keys['other@example.org'], 1)
        self.admin()
        raw = dict(name='Вся организация', period='range', date_from='2026-09-07', date_to='2026-09-13')
        page = self.client.post('/manager/filters', data=raw)
        url = self.export_link(page)
        _, rows = self.exported_rows(self.client.get(url))
        self.assertEqual([r['ID заявки'] for r in rows], [1])
        with ui.get_db_connection() as conn:
            saved = conn.execute('SELECT owner_key,payload FROM app_team_filter WHERE name=?', (raw['name'],)).fetchone()
            self.assertEqual(saved['owner_key'], d.SUPERUSER_KEY)
            self.assertEqual(json.loads(saved['payload'])['scope'], 'all')
        with TestClient(ui.app) as other:
            other.post('/superuser/login', data={'superuser_login':'root', 'superuser_password':'test-password'})
            self.assertIn('Уникальная задача', other.get('/manager', params={'saved':raw['name']}).text)
            other.post('/manager/filters/delete', data={'name':raw['name']})
        self.assertEqual(self.client.get('/manager', params={'saved':raw['name']}).status_code, 404)
        self.client.post('/superuser/logout')
        for path in ('/manager', '/manager/requests/1', url):
            self.assertEqual(self.client.get(path).status_code, 403)
        self.login()
        self.assertEqual(self.client.get('/manager').status_code, 200)
        self.assertEqual(self.client.get('/manager/requests/1').status_code, 404)
        self.assertEqual(self.client.get('/manager', params={'saved':raw['name']}).status_code, 404)

    def test_superuser_invalid_session_and_employee_admin_do_not_get_global_access(self):
        other_key = self.keys['other@example.org']
        self.task(other_key, 1)
        self.client.cookies.set(ui.SUPERUSER_COOKIE_NAME, '__superuser__')
        self.assertEqual(self.client.get('/manager').status_code, 403)
        self.login('worker@example.org')
        with ui.get_db_connection() as conn:
            ui.update_employee_admin_role(conn, self.worker, True, 'test')
        self.assertEqual(self.client.get('/admin/users').status_code, 200)
        for path in ('/manager', '/manager/requests/1', '/manager/requests/export'):
            self.assertEqual(self.client.get(path, params={'owner_key':d.SUPERUSER_KEY}).status_code, 403)
        self.admin()
        url = self.export_link(self.client.get('/manager'))
        with patch.dict('os.environ', {'WORK_ON_HOLIDAY_SUPERUSER_PASSWORD':'changed-password'}):
            self.assertEqual(self.client.get(url).status_code, 403)

    def test_superuser_roster_preview_default_manager_and_admin_mutation(self):
        self.admin()
        data = (';'.join(d.HEADERS) + '\nНовый Сотрудник;new@example.org;').encode()
        preview = self.client.post('/admin/roster/preview', files={'file':('root.csv', data)})
        self.assertIn('<td>Суперпользователь</td>', preview.text)
        identifier = re.search(r'/admin/roster/([^/]+)/apply', preview.text)[1]
        self.client.post(f'/admin/roster/{identifier}/apply')
        with ui.get_db_connection() as conn:
            people = d.employees(conn)
            new = next(p for p in people.values() if p['email'] == 'new@example.org')
            self.assertEqual(new['manager'], d.SUPERUSER_KEY)
        self.task(self.worker, 1)
        detail = self.client.get('/manager/requests/1')
        link = unescape(re.search(r'href="([^"]+)">Управлять заявками', detail.text)[1])
        self.assertIn('Уникальная задача', self.client.get(link).text)
        self.client.post('/employee/request/cancel', data={'employee_key':self.worker, 'response_id':1, 'admin_mode':'1'})
        with ui.get_db_connection() as conn:
            self.assertEqual(conn.execute('SELECT status FROM app_request_state WHERE response_id=1').fetchone()[0], 'cancelled')

    def test_manager_export_matches_filtered_page_and_corrected_values(self):
        self.login()
        direct = self.keys['direct@example.org']
        other = self.keys['other@example.org']
        self.task(self.worker, 1, day='2026-08-01', text='Старое описание')
        self.task(direct, 2, day='2026-09-13', text='Релиз прямой')
        self.task(other, 3, text='Релиз чужой')
        self.task(self.worker, 4, day='2026-09-14', text='Релиз за периодом')
        self.task(self.worker, 5, text='Другой текст')
        with ui.get_db_connection() as conn:
            conn.execute('''INSERT INTO app_request_state(request_uid,response_id,full_name_key,status,
                override_planned_work_date,override_task_description,override_justification,override_systems,
                actual_work_date,actual_work_time,created_at,updated_at) VALUES ('req:1',1,?,'cancelled','2026-09-07',
                'Релиз исправленный','Новое обоснование','Система 1 | Система 2','2026-09-08','11:00 - 13:00',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)''', (self.worker,))
        params = dict(apply='1', period='range', date_from='2026-09-07', date_to='2026-09-13',
                      scope='all', employees=[self.worker, direct, other], surname='сОТРУДНИК', text='РЕЛИЗ')
        page = self.client.get('/manager', params=params)
        book, rows = self.exported_rows(self.client.get(self.export_link(page)))
        self.assertEqual([r['ID заявки'] for r in rows],
                         [int(value) for value in re.findall(r'href="/manager/requests/(\d+)"', page.text)])
        self.assertEqual({r['ID заявки'] for r in rows}, {1, 2})
        corrected = next(r for r in rows if r['ID заявки'] == 1)
        self.assertEqual(corrected['Плановая дата'].date(), date(2026, 9, 7))
        self.assertEqual(corrected['Задача'], 'Релиз исправленный')
        self.assertEqual(corrected['Статус'], 'Отменена')
        self.assertEqual(corrected['Обоснование'], 'Новое обоснование')
        self.assertEqual(corrected['Информационные системы'], 'Система 1 | Система 2')
        self.assertEqual(corrected['Фактическая дата'].date(), date(2026, 9, 8))
        self.assertEqual(corrected['Фактическое время'], '11:00 - 13:00')
        self.assertIn('Начальник Средний', corrected['Руководитель'])
        self.assertEqual(book['Заявки'].freeze_panes, 'C2')
        self.assertEqual(dict(book['Фильтры'].values)['Количество заявок'], 2)
        params['scope'] = 'direct'
        _, rows = self.exported_rows(self.client.get('/manager/requests/export', params=params))
        self.assertEqual([r['ID заявки'] for r in rows], [2])

    def test_manager_download_keeps_applied_filter_and_does_not_change_last_filter(self):
        self.login()
        self.task(self.worker, 1)
        self.task(self.worker, 2, day='2026-09-19')
        raw = dict(name='Сохранённый', period='range', scope='all', date_from='2026-09-07',
                   date_to='2026-09-13', employees=[self.worker])
        page = self.client.post('/manager/filters', data=raw)
        url = self.export_link(page)
        # Another tab changes the stored filter; the old page must still export its selection.
        self.client.post('/manager/filters', data={**raw, 'date_from':'2026-09-14', 'date_to':'2026-09-20'})
        with ui.get_db_connection() as conn:
            before = conn.execute('SELECT name,payload FROM app_team_filter WHERE owner_key=? ORDER BY name', (self.root,)).fetchall()
            before = [tuple(row) for row in before]
        _, rows = self.exported_rows(self.client.get(url))
        self.assertEqual([r['ID заявки'] for r in rows], [1])
        with ui.get_db_connection() as conn:
            after = [tuple(row) for row in conn.execute('SELECT name,payload FROM app_team_filter WHERE owner_key=? ORDER BY name', (self.root,))]
        self.assertEqual(before, after)
        _, rows = self.exported_rows(self.client.get('/manager/requests/export', params={'saved':'Сохранённый'}))
        self.assertEqual([r['ID заявки'] for r in rows], [2])
        page = self.client.get('/manager', params={'apply':'1', 'period':'current'})
        query = parse_qs(urlsplit(self.export_link(page)).query)
        self.assertEqual(query['period'], ['range'])
        self.assertIn('date_from', query)
        self.assertIn('date_to', query)

    def test_manager_export_rechecks_access_and_stale_selection(self):
        self.login()
        self.task(self.worker, 1)
        self.task(self.keys['direct@example.org'], 2)
        self.task(self.keys['other@example.org'], 3)
        params = dict(apply='1', period='range', scope='all', date_from='2026-09-07',
                      date_to='2026-09-13', employees=[self.worker])
        url = self.export_link(self.client.get('/manager', params=params))
        with ui.get_db_connection() as conn:
            d.apply_import(conn, d.plan_import(conn, [['Сотрудник Нижний', 'worker@example.org', 'other@example.org']]))
        book, rows = self.exported_rows(self.client.get(url))
        self.assertEqual(rows, [])
        self.assertEqual(dict(book['Фильтры'].values)['Количество заявок'], 0)
        # Selecting unrelated identities directly cannot expand permissions.
        _, rows = self.exported_rows(self.client.get('/manager/requests/export', params={**params, 'employees':[self.worker, self.keys['other@example.org']]}))
        self.assertEqual(rows, [])
        with ui.get_db_connection() as conn:
            d.apply_import(conn, d.plan_import(conn, [['Начальник Средний', 'middle@example.org', ''], ['Сотрудник Прямой', 'direct@example.org', '']]))
        self.assertEqual(self.client.get(url).status_code, 403)

    def test_export_requires_current_role(self):
        for url in ['/admin/requests/export', '/manager/requests/export']:
            self.assertEqual(self.client.get(url).status_code, 403)
        self.login('worker@example.org')
        for url in ['/admin/requests/export', '/manager/requests/export']:
            self.assertEqual(self.client.get(url).status_code, 403)
        self.login()
        self.assertEqual(self.client.get('/admin/requests/export').status_code, 403)
        with ui.get_db_connection() as conn:
            ui.update_employee_admin_role(conn, self.root, True, 'test')
        self.exported_rows(self.client.get('/admin/requests/export'))
        with ui.get_db_connection() as conn:
            ui.update_employee_status(conn, self.root, 'blocked', 'test', 'test')
        for url in ['/admin/requests/export', '/manager/requests/export']:
            self.assertEqual(self.client.get(url).status_code, 403)

    def test_admin_export_combines_live_filters_and_preserves_all_statuses(self):
        self.admin()
        self.task(self.worker, 1, day='2026-08-01')
        self.task(self.worker, 2, day='2026-09-13')
        self.task(self.keys['direct@example.org'], 3)
        with ui.get_db_connection() as conn:
            conn.execute('''INSERT INTO app_request_state(request_uid,response_id,full_name_key,status,
                override_planned_work_date,override_task_description,override_justification,override_systems,
                actual_work_date,actual_work_time,created_at,updated_at) VALUES ('req:1',1,?,'cancelled','2026-09-12',
                'Исправленный текст','Причина','Новая система','2026-09-13','12:00 - 14:00',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)''', (self.worker,))
        params = dict(filter_name='  НИЖНИЙ ', filter_status='cancelled', filter_date='12/09/2026')
        page = self.client.get('/admin/requests', params=params)
        _, rows = self.exported_rows(self.client.get(self.export_link(page)))
        self.assertEqual([r['ID заявки'] for r in rows], [1])
        self.assertEqual(rows[0]['Задача'], 'Исправленный текст')
        self.assertEqual(rows[0]['Обоснование'], 'Причина')
        self.assertEqual(rows[0]['Информационные системы'], 'Новая система')
        self.assertEqual(rows[0]['Плановая дата'].date(), date(2026, 9, 12))
        self.assertEqual(rows[0]['Фактическое время'], '12:00 - 14:00')
        cases = [({}, [2, 1, 3]), ({'filter_name':'нижний'}, [2, 1]),
                 ({'filter_date':'12/09/2026'}, [1, 3]), ({'filter_status':'active'}, [2, 3]),
                 ({'filter_date':'12/09'}, [2, 1, 3]), ({'filter_date':'31/02/2026'}, []),
                 ({'filter_name':'Не существует'}, [])]
        for query, expected in cases:
            with self.subTest(query=query):
                _, rows = self.exported_rows(self.client.get('/admin/requests/export', params=query))
                self.assertEqual([r['ID заявки'] for r in rows], expected)

    def test_export_preserves_literal_formula_text_and_invalid_filters_fail(self):
        self.login()
        text = '=HYPERLINK("https://example.org", "Текст задачи")'
        self.task(self.worker, 1, text=text)
        params = dict(apply='1', scope='all', period='range', date_from='2026-09-07', date_to='2026-09-13')
        book, rows = self.exported_rows(self.client.get('/manager/requests/export', params=params))
        self.assertEqual(rows[0]['Задача'], text)
        cell = next(cell for row in book['Заявки'] for cell in row if cell.value == text)
        self.assertEqual(cell.data_type, 's')
        for query in [{**params, 'date_to':'2026-09-01'}, {**params, 'scope':'invalid'}, {'saved':'Unknown'}]:
            self.assertIn(self.client.get('/manager/requests/export', params=query).status_code, {404, 422})

    def test_migration_idempotent_preserves_legacy(self):
        with ui.get_db_connection() as conn:
            legacy = ui.register_employee_directory_entry(conn, 'Старый Сотрудник')['employee_key']
            ui.upsert_employee_token(conn, legacy, 'secret')
            ui.update_employee_admin_role(conn, legacy, True, 'test')
            ensure_app_tables(conn)
            ensure_app_tables(conn)
            plan = d.plan_import(conn, [['Старый Сотрудник','legacy@example.org','root@example.org']])
            self.assertEqual(plan['skipped'], [])
            self.assertEqual(plan['changes'][0]['key'], legacy)
            d.apply_import(conn, plan)
            self.assertEqual(ui.get_employee_profile(conn, legacy)['is_admin'], 1)
            self.assertEqual(ui.get_employee_token_record(conn, legacy)['token_hash'], ui.hash_employee_token('secret'))
            self.assertIn(legacy, d.team_keys(conn, self.root))

    def test_hierarchy_all_depths_and_default(self):
        with ui.get_db_connection() as conn:
            self.assertEqual(len(d.team_keys(conn, self.root)), 2)
            self.assertEqual(len(d.team_keys(conn, self.root, 'all')), 3)
            extra = [[f'Уровень {i}', f'level{i}@example.org', 'worker@example.org' if i == 0 else f'level{i-1}@example.org'] for i in range(30)]
            d.apply_import(conn, d.plan_import(conn, extra))
            self.assertEqual(len(d.team_keys(conn, self.root, 'all')), 33)

    def test_cycles_duplicates_unknown_and_atomicity(self):
        bad_sets = [ROWS + [ROWS[0]], [['Руководитель Главный','root@example.org','worker@example.org']],
                    [['Некто','new@example.org','unknown@example.org']], [['Некто','new@example.org','new@example.org']],
                    [['Некто','invalid','']]]
        with ui.get_db_connection() as conn:
            before = d.fingerprint(conn)
            for rows in bad_sets:
                plan = d.plan_import(conn, rows)
                self.assertTrue(plan['errors'])
                with self.assertRaises(ValueError): d.apply_import(conn, plan)
                self.assertEqual(d.fingerprint(conn), before)

    def test_omitted_unchanged_and_empty_manager(self):
        with ui.get_db_connection() as conn:
            before = d.employees(conn)[self.worker]
            d.apply_import(conn, d.plan_import(conn, [['Начальник Средний','middle@example.org','']]))
            self.assertEqual(d.employees(conn)[self.worker], before)
            self.assertNotIn(self.worker, d.team_keys(conn, self.root, 'all'))

    def test_namesakes_distinct_and_email_login(self):
        with ui.get_db_connection() as conn:
            d.apply_import(conn, d.plan_import(conn, [['Сотрудник Нижний','twin@example.org','']]))
            with self.assertRaises(Exception): ui.resolve_employee_by_name(conn, 'Сотрудник Нижний')
            self.assertEqual(ui.resolve_employee_by_name(conn, 'WORKER@example.org')['employee_key'], self.worker)
        response = self.client.post('/employee/login', data={'full_name':'worker@example.org'}, follow_redirects=False)
        self.assertIn('worker%40example.org', response.headers['location'])
        response = self.client.post('/employee/login', data={'full_name':'worker@example.org','access_token':'worker@example.org'})
        self.assertIn('Сотрудник Нижний', response.text)

    def test_imported_first_token_admin_only(self):
        with ui.get_db_connection() as conn:
            conn.execute('DELETE FROM app_employee_auth WHERE full_name_key=?', (self.root,))
        response = self.client.post('/employee/login', data={'full_name':'root@example.org'})
        self.assertIn('получите персональный токен', response.text)
        with ui.get_db_connection() as conn:
            self.assertIsNone(ui.get_employee_token_record(conn, self.root))

    def test_csv_excel(self):
        csv_data = (';'.join(d.HEADERS) + '\n' + '\n'.join(';'.join(r) for r in ROWS)).encode('utf-8-sig')
        self.assertEqual(d.parse_roster('data.csv', csv_data), ROWS)
        book = Workbook()
        for row in [d.HEADERS] + ROWS: book.active.append(row)
        stream = io.BytesIO(); book.save(stream)
        self.assertEqual(d.parse_roster('data.xlsx', stream.getvalue()), ROWS)
        with self.assertRaises(ValueError): d.parse_roster('data.xls', b'bad')

    def test_preview_apply_permissions_and_stale(self):
        data = (';'.join(d.HEADERS)+'\nНовый Сотрудник;new@example.org;root@example.org').encode()
        self.assertEqual(self.client.post('/admin/roster/preview',files={'file':('x.csv',data)}).status_code,403)
        self.admin()
        response = self.client.post('/admin/roster/preview',files={'file':('x.csv',data)})
        self.assertEqual(response.status_code,200)
        identifier = re.search(r'/admin/roster/([^/]+)/apply', response.text)[1]
        with ui.get_db_connection() as conn:
            self.assertNotIn('new@example.org', [p['email'] for p in d.employees(conn).values()])
            d.apply_import(conn, d.plan_import(conn, [['Сторонний','outsider@example.org','']]))
        self.assertEqual(self.client.post(f'/admin/roster/{identifier}/apply').status_code,409)
        response = self.client.post('/admin/roster/preview',files={'file':('x.csv',data)})
        identifier = re.search(r'/admin/roster/([^/]+)/apply', response.text)[1]
        self.assertEqual(self.client.post(f'/admin/roster/{identifier}/apply',follow_redirects=False).status_code,303)
        self.assertEqual(self.client.post(f'/admin/roster/{identifier}/apply').status_code,404)

    def test_list_filters_details_and_no_work_inference(self):
        self.login()
        self.task(self.worker,1)
        self.task(self.keys['direct@example.org'],2,text='Прямая задача')
        self.task(self.keys['other@example.org'],3,text='Чужой секрет')
        params = dict(apply='1',period='range',date_from='2026-09-07',date_to='2026-09-13')
        response = self.client.get('/manager',params=params)
        self.assertNotIn('Уникальная задача',response.text)
        self.assertIn('Прямая задача',response.text)
        response = self.client.get('/manager',params={**params,'scope':'all','text':'Уникальная'})
        self.assertIn('Уникальная задача',response.text)
        self.assertNotIn('Чужой секрет',response.text)
        self.assertIn('Сотрудники без задач в периоде · 1',response.text)
        self.assertEqual(self.client.get('/manager/requests/1').status_code,200)
        self.assertEqual(self.client.get('/manager/requests/3').status_code,404)
        self.assertEqual(self.client.get('/manager/requests/999').status_code,404)
        self.assertNotIn('/employee/request/correct',self.client.get('/manager/requests/1').text)

    def test_saved_filters_cross_client_and_scope_revocation(self):
        self.login()
        self.task(self.worker,1)
        raw = dict(name='Моя команда',scope='all',period='range',date_from='2026-09-07',date_to='2026-09-13',employees=[self.worker])
        response = self.client.post('/manager/filters',data=raw)
        self.assertIn('Уникальная задача',response.text)
        with TestClient(ui.app) as other:
            other.cookies.set(ui.EMPLOYEE_TOKEN_COOKIE_NAME,'root@example.org')
            self.assertIn('Уникальная задача',other.get('/manager').text)
            other.cookies.set(ui.EMPLOYEE_TOKEN_COOKIE_NAME,'middle@example.org')
            self.assertEqual(other.get('/manager',params={'saved':'Моя команда'}).status_code,404)
        with ui.get_db_connection() as conn:
            d.apply_import(conn,d.plan_import(conn,[['Сотрудник Нижний','worker@example.org','other@example.org']]))
        self.assertNotIn('Уникальная задача',self.client.get('/manager',params={'saved':'Моя команда'}).text)
        self.assertEqual(self.client.get('/manager/requests/1').status_code,404)

    def test_current_week_relative_and_navigation(self):
        filters=d.normalize_filters({'period':'current'})
        self.assertEqual(d.filter_dates(filters,date(2026,9,9)),('2026-09-07','2026-09-13'))
        self.assertEqual(d.filter_dates(filters,date(2026,9,14)),('2026-09-14','2026-09-20'))
        self.login()
        response=self.client.get('/manager',params={'apply':'1','period':'week','date_from':'2026-09-09','step':'1'})
        self.assertIn('2026-09-14 — 2026-09-20',response.text)
        self.assertEqual(self.client.get('/manager',params={'apply':'1','period':'range','date_from':'2026-09-09','date_to':'2026-09-01'}).status_code,422)

    def test_manager_role_lost_but_admin_retained(self):
        self.login()
        with ui.get_db_connection() as conn:
            ui.update_employee_admin_role(conn,self.root,True,'test')
            d.apply_import(conn,d.plan_import(conn,[['Начальник Средний','middle@example.org',''],['Сотрудник Прямой','direct@example.org','']]))
        self.assertEqual(self.client.get('/manager').status_code,403)
        self.assertEqual(self.client.get('/admin/users').status_code,200)
        self.assertEqual(self.client.get('/employee').status_code,200)

    def test_blocked_manager_and_nonmanager_denied(self):
        self.login('worker@example.org')
        self.assertEqual(self.client.get('/manager').status_code,403)
        self.login()
        with ui.get_db_connection() as conn:
            ui.update_employee_status(conn,self.root,'blocked','test','test')
        self.assertEqual(self.client.get('/manager').status_code,403)

    def test_manager_cannot_mutate_subordinate(self):
        self.login()
        self.task(self.worker,1)
        response=self.client.post('/employee/request/cancel',data={'employee_key':self.worker,'response_id':1})
        with ui.get_db_connection() as conn:
            row=conn.execute("SELECT status FROM app_request_state WHERE response_id=1").fetchone()
            self.assertTrue(row is None or row[0]!='cancelled')

    def test_preview_automatically_links_legacy_name_and_preserves_token(self):
        with ui.get_db_connection() as conn:
            legacy=ui.register_employee_directory_entry(conn,'Старый Сотрудник')['employee_key']
            ui.upsert_employee_token(conn,legacy,'legacy-token')
        self.admin()
        content=(';'.join(d.HEADERS)+'\nСтарый Сотрудник;legacy@example.org;root@example.org').encode()
        response=self.client.post('/admin/roster/preview',files={'file':('x.csv',content)})
        identifier=re.search(r'/admin/roster/([^/]+)/apply',response.text)[1]
        self.assertIn('ФИО — почта будет заполнена', response.text)
        with ui.get_db_connection() as conn:
            self.assertEqual(d.employees(conn)[legacy]['email'], '')
        self.assertIn('Применить записи (1)',response.text)
        self.assertEqual(self.client.post(f'/admin/roster/{identifier}/apply',follow_redirects=False).status_code,303)
        with ui.get_db_connection() as conn:
            self.assertEqual(d.employees(conn)[legacy]['email'],'legacy@example.org')
            self.assertEqual(ui.get_employee_token_record(conn,legacy)['token_hash'],ui.hash_employee_token('legacy-token'))

    def test_effective_corrected_dates_text_and_multiselect(self):
        self.task(self.worker,1,day='2026-08-01',text='Старое описание')
        direct=self.keys['direct@example.org']
        self.task(direct,2,text='Прямая задача')
        with ui.get_db_connection() as conn:
            conn.execute('''INSERT INTO app_request_state(request_uid,response_id,full_name_key,status,
                override_planned_work_date,override_task_description,created_at,updated_at)
                VALUES ('req:1',1,?,'active','2026-09-12','Исправленная задача',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)''',(self.worker,))
        self.login()
        params={'apply':'1','period':'range','scope':'all','date_from':'2026-09-07','date_to':'2026-09-13','employees':[self.worker,direct]}
        response=self.client.get('/manager',params=params)
        self.assertIn('Исправленная задача',response.text)
        self.assertIn('Прямая задача',response.text)
        response=self.client.get('/manager',params={**params,'surname':'Нижний','text':'Исправленная'})
        self.assertIn('Исправленная задача',response.text)
        self.assertNotIn('Прямая задача',response.text)
        self.assertIn('Исправленная задача',self.client.get('/manager/requests/1').text)

    def test_manager_creates_own_request(self):
        self.login()
        response=self.client.post('/employee/request/create',data={'planned_work_date':'2099-09-12',
            'planned_work_time':'10:00 - 12:00','task_description':'Собственная задача','justification':'Проверка',
            'systems':'Тестовая АС','payment_type':'Отгул'})
        with ui.get_db_connection() as conn:
            row=conn.execute('SELECT full_name_key FROM survey_responses WHERE task_description=?',('Собственная задача',)).fetchone()
            self.assertIsNotNone(row,response.text)
            self.assertEqual(row[0],self.root)

    def test_saved_overwrite_delete_and_last_filter(self):
        self.login()
        self.client.post('/manager/filters',data={'name':'Обзор','scope':'all','period':'current'})
        self.client.post('/manager/filters',data={'name':'Обзор','scope':'direct','period':'current','text':'новое'})
        with ui.get_db_connection() as conn:
            rows=conn.execute("SELECT payload FROM app_team_filter WHERE owner_key=? AND name='Обзор'",(self.root,)).fetchall()
            self.assertEqual(len(rows),1)
            self.assertEqual(json.loads(rows[0][0])['text'],'новое')
        self.client.post('/manager/filters/delete',data={'name':'Обзор'})
        self.assertEqual(self.client.get('/manager',params={'saved':'Обзор'}).status_code,404)
        self.assertIn('value="новое"',self.client.get('/manager').text)

    def test_existing_email_and_duplicate_legacy_email(self):
        with ui.get_db_connection() as conn:
            legacy=ui.register_employee_directory_entry(conn,'Внешний Руководитель')['employee_key']
            conn.execute('UPDATE app_employee_directory SET work_email=? WHERE full_name_key=?',('external@example.org',legacy))
            plan=d.plan_import(conn,[['Сотрудник Нижний','WORKER@example.org','external@example.org']])
            self.assertFalse(plan['errors'])
            d.apply_import(conn,plan)
            self.assertIn(self.worker,d.team_keys(conn,legacy))
            duplicate=ui.register_employee_directory_entry(conn,'Другая Запись')['employee_key']
            conn.execute('UPDATE app_employee_directory SET work_email=? WHERE full_name_key=?',('external@example.org',duplicate))
            self.assertTrue(d.plan_import(conn,ROWS)['errors'])

    def test_legacy_rename_cannot_be_overwritten_by_guest(self):
        with ui.get_db_connection() as conn:
            key=ui.register_employee_directory_entry(conn,'Прежнее Имя')['employee_key']
            d.apply_import(conn,d.plan_import(conn,[['Прежнее Имя','rename@example.org','']]))
            d.apply_import(conn,d.plan_import(conn,[['Новое Имя','rename@example.org','']]))
        self.client.post('/employee/login',data={'full_name':'Прежнее Имя'})
        with ui.get_db_connection() as conn:
            self.assertEqual(d.employees(conn)[key]['name'],'Новое Имя')
