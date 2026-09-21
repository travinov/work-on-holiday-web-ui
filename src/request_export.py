"""Private XLSX downloads of the current request selection."""
from datetime import date, datetime
from io import BytesIO
from urllib.parse import quote

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def request_export_response(rows, *, filename, filters, include_manager=False):
    columns = [("response_id", "ID заявки", 12), ("full_name", "ФИО", 32)]
    if include_manager:
        columns.append(("manager_name", "Руководитель", 38))
    columns += [
        ("planned_work_date", "Плановая дата", 16),
        ("planned_work_time", "Плановое время", 20),
        ("status_label", "Статус", 32),
        ("payment_type", "Условия выхода", 24),
        ("task_description", "Задача", 55),
        ("justification", "Обоснование", 42),
        ("systems", "Информационные системы", 36),
        ("actual_work_date", "Фактическая дата", 18),
        ("actual_work_time", "Фактическое время", 20),
        ("planning_lock_label", "Закрытие приёма заявок", 26),
        ("actual_lock_label", "Закрытие ввода факта", 26),
        ("lock_week_label", "Ранее закрытая неделя", 26),
    ]
    book = Workbook()
    sheet = book.active
    sheet.title = "Заявки"

    def put(sheet, row, column, value):
        cell = sheet.cell(row, column)
        if isinstance(value, str):
            # User text must stay text, including strings starting with '='.
            cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
            cell.data_type = "s"
        else:
            cell.value = value
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if isinstance(value, date):
            cell.number_format = "DD.MM.YYYY"
        return cell

    for column, (_, title, width) in enumerate(columns, 1):
        cell = put(sheet, 1, column, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 32
    for row_number, item in enumerate(rows, 2):
        for column, (key, _, _) in enumerate(columns, 1):
            value = item.get(key) or item.get(key + "_iso") or ""
            if key in {"planned_work_date", "actual_work_date"} and value:
                try:
                    value = date.fromisoformat(value)
                except ValueError:
                    pass
            cell = put(sheet, row_number, column, value)
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EFF6FC")
        sheet.row_dimensions[row_number].height = 60
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions

    metadata = book.create_sheet("Фильтры")
    details = [
        ("Реестр", "Отфильтрованный список заявок"),
        ("Сформирован (время сервера)", datetime.now().isoformat(timespec="seconds")),
        ("Количество заявок", len(rows)),
        *filters,
        ("Примечание", "Данные и права доступа проверены на момент выгрузки. Даты фильтра относятся к плановому выходу."),
    ]
    for row_number, (label, value) in enumerate(details, 1):
        put(metadata, row_number, 1, label).font = Font(bold=True)
        put(metadata, row_number, 2, value)
    metadata.column_dimensions["A"].width = 34
    metadata.column_dimensions["B"].width = 85
    metadata.freeze_panes = "B2"
    output = BytesIO()
    book.save(output)
    book.close()
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename, safe='')}",
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )
