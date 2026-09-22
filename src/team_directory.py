"""Email identities, atomic roster imports and hierarchy traversal.

Legacy employee keys remain stable, preserving tokens, profiles and request history.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import uuid
from datetime import date, timedelta


HEADERS = ['ФИО сотрудника', 'Рабочая почта сотрудника', 'Рабочая почта руководителя']
SUPERUSER_KEY = '__superuser__'
SUPERUSER_NAME = 'Суперпользователь'


def ensure_team_tables(conn):
    conn.execute('''CREATE TABLE IF NOT EXISTS app_org_employee (
        employee_key TEXT PRIMARY KEY, email TEXT NOT NULL UNIQUE COLLATE NOCASE,
        manager_key TEXT, CHECK(employee_key != manager_key))''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_org_manager ON app_org_employee(manager_key)')
    conn.execute('''CREATE TABLE IF NOT EXISTS app_team_filter (
        owner_key TEXT NOT NULL, name TEXT NOT NULL, payload TEXT NOT NULL,
        PRIMARY KEY(owner_key, name))''')
    conn.execute('''CREATE TABLE IF NOT EXISTS app_roster_draft (
        id TEXT PRIMARY KEY, actor TEXT NOT NULL, payload TEXT NOT NULL,
        bindings TEXT NOT NULL, fingerprint TEXT NOT NULL,
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)''')


def name_key(value):
    return ' '.join((value or '').lower().replace('ё', 'е').split())


def email(value):
    value = str(value or '').strip().lower()
    if value and not re.fullmatch(r'[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+', value):
        raise ValueError('Некорректная рабочая почта: ' + value)
    return value


def parse_roster(filename, data):
    if len(data) > 5 * 1024 * 1024:
        raise ValueError('Размер файла не должен превышать 5 МБ')
    if filename.lower().endswith('.csv'):
        try:
            content = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            content = data.decode('cp1251')
        try:
            dialect = csv.Sniffer().sniff(content[:4096], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(content), dialect))
    elif filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        from zipfile import ZipFile
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError('Слишком большой распакованный Excel')
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            rows = []
            for row in book.active.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > 10001:
                    raise ValueError('Не более 10000 сотрудников в одном файле')
        finally:
            book.close()
    else:
        raise ValueError('Поддерживаются CSV и Excel .xlsx')
    rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
    if not rows or [str(v or '').strip() for v in rows[0][:3]] != HEADERS:
        raise ValueError('Заголовки должны соответствовать шаблону')
    if len(rows) < 2 or len(rows) > 10001:
        raise ValueError('В файле должно быть от 1 до 10000 сотрудников')
    return [[str(v or '').strip() for v in (r + [''] * 3)[:3]] for r in rows[1:]]


def employees(conn):
    result = {}
    for r in conn.execute('''SELECT full_name_key, COALESCE(full_name_normalized,full_name) AS full_name
                            FROM survey_responses WHERE request_type='Подать заявку' AND full_name_key IS NOT NULL
                            ORDER BY response_id'''):
        result[r['full_name_key']] = dict(key=r['full_name_key'], name=r['full_name'], email='', manager=None)
    for r in conn.execute('SELECT * FROM app_employee_directory'):
        result[r['full_name_key']] = dict(key=r['full_name_key'], name=r['full_name'],
                                         email=(r['work_email'] or '').strip().lower(), manager=None)
    for r in conn.execute('SELECT * FROM app_org_employee'):
        if r['employee_key'] in result:
            result[r['employee_key']].update(email=r['email'], manager=r['manager_key'])
    # The root is a system account, never an imported employee or a token identity.
    # Include legacy and self-registered employees without rewriting their records.
    result.pop(SUPERUSER_KEY, None)
    for person in result.values():
        if person['manager'] not in result:
            person['manager'] = SUPERUSER_KEY
    return result


def manager_name(people, key):
    if key == SUPERUSER_KEY:
        return SUPERUSER_NAME
    return people.get(key, {}).get('name', 'Без руководителя')


def fingerprint(conn):
    return hashlib.sha256(json.dumps(employees(conn), sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def plan_import(conn, rows):
    current = employees(conn)
    by_email, by_name = {}, {}
    errors, changes, skipped = [], [], []
    for key, person in current.items():
        by_name.setdefault(name_key(person['name']), []).append(person)
        if person['email']:
            by_email.setdefault(person['email'], []).append(key)
    for address, keys in by_email.items():
        if len(keys) > 1:
            errors.append('В базе несколько сотрудников с почтой ' + address + '. Исправьте справочник до импорта.')
    seen = set()
    graph = {k: p['manager'] for k, p in current.items()}
    current_email_keys = {a: ks[0] for a, ks in by_email.items() if len(ks) == 1}

    def skip(row, reason):
        skipped.append(dict(row=row['row'], name=row['name'], email=row['email'], reason=reason))

    for index, row in enumerate(rows, 2):
        name, address, manager = row
        try:
            address, manager = email(address), email(manager)
            if not name or not address:
                raise ValueError('ФИО и почта сотрудника обязательны')
            if len(name) > 250 or len(address) > 254 or len(manager) > 254:
                raise ValueError('Слишком длинное ФИО или почта')
            if address in seen:
                raise ValueError('Дублирующаяся почта в файле: ' + address)
            seen.add(address)
            key = current_email_keys.get(address)
            matched_by = 'Рабочая почта'
            if not key:
                namesakes = by_name.get(name_key(name), [])
                candidates = [p for p in namesakes if not p['email']]
                if candidates and len(namesakes) > 1:
                    skip(dict(row=index, name=name, email=address),
                         'Неоднозначное ФИО: в системе несколько сотрудников с таким именем.')
                    continue
                if candidates:
                    key = candidates[0]['key']
                    matched_by = 'ФИО — почта будет заполнена'
                else:
                    key = 'emp:' + str(uuid.uuid5(uuid.NAMESPACE_URL, 'woh:' + address))
                    matched_by = 'Новая запись'
            changes.append(dict(row=index, key=key, name=name, email=address, manager_email=manager,
                                matched_by=matched_by,
                                old=current.get(key), action='Обновление' if key in current else 'Создание'))
        except ValueError as exc:
            errors.append(f'Строка {index}: {exc}')

    # Decide all name matches together: file order must not choose which email
    # wins when several rows refer to the same legacy employee.
    claims = {}
    for item in changes:
        claims.setdefault(item['key'], []).append(item)
    ambiguous = {key for key, items in claims.items() if len(items) > 1}
    for item in changes:
        if item['key'] in ambiguous:
            skip(item, 'Несколько строк файла сопоставлены одному сотруднику. Уточните ФИО и почту.')
    changes = [item for item in changes if item['key'] not in ambiguous]

    # A skipped new manager has no usable identity. Skip dependent rows too,
    # but an existing manager remains usable through their unchanged record.
    unavailable = {item['email'] for item in skipped} - current_email_keys.keys()
    pending, dependent_rows, by_manager = list(unavailable), set(), {}
    for item in changes:
        by_manager.setdefault(item['manager_email'], []).append(item)
    for manager_address in pending:
        for item in by_manager.get(manager_address, []):
            if item['row'] in dependent_rows:
                continue
            skip(item, 'Руководитель пропущен при сопоставлении: ' + item['manager_email'])
            dependent_rows.add(item['row'])
            if item['email'] not in current_email_keys and item['email'] not in unavailable:
                unavailable.add(item['email'])
                pending.append(item['email'])
    changes = [item for item in changes if item['row'] not in dependent_rows]

    email_keys = {**current_email_keys, **{item['email']: item['key'] for item in changes}}
    for item in changes:
        manager = email_keys.get(item['manager_email']) if item['manager_email'] else SUPERUSER_KEY
        if item['manager_email'] and manager is None:
            errors.append('Неизвестный руководитель: ' + item['manager_email'])
        item['manager'] = manager
        graph[item['key']] = manager
    for start in graph:
        trail, node = set(), start
        while node is not None:
            if node in trail:
                errors.append('Цикл подчинения: ' + start)
                break
            trail.add(node)
            node = graph.get(node)
        if errors and errors[-1].startswith('Цикл'):
            break
    return dict(changes=changes, errors=errors, skipped=sorted(skipped, key=lambda item: item['row']))


def apply_import(conn, plan):
    if plan['errors']:
        raise ValueError('Сначала устраните ошибки в реестре')
    if not plan['changes']:
        return
    # Materialize known email identities so a manager outside the file is usable.
    for person in employees(conn).values():
        if person['email']:
            conn.execute('INSERT OR IGNORE INTO app_org_employee(employee_key,email,manager_key) VALUES (?,?,?)',
                         (person['key'], person['email'], person['manager']))
    for p in plan['changes']:
        conn.execute('''INSERT INTO app_employee_directory(full_name_key,full_name,work_email,created_at,updated_at)
                        VALUES (?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)
                        ON CONFLICT(full_name_key) DO UPDATE SET full_name=excluded.full_name,
                        work_email=excluded.work_email,updated_at=CURRENT_TIMESTAMP''', (p['key'], p['name'], p['email']))
        conn.execute('''INSERT INTO app_org_employee(employee_key,email,manager_key) VALUES (?,?,?)
                        ON CONFLICT(employee_key) DO UPDATE SET email=excluded.email,manager_key=excluded.manager_key''',
                     (p['key'], p['email'], p['manager']))


def team_keys(conn, owner, scope='direct'):
    if owner == SUPERUSER_KEY:
        people = employees(conn)
        return {key for key, person in people.items()
                if scope == 'all' or person['manager'] == SUPERUSER_KEY}
    if scope == 'direct':
        return {r[0] for r in conn.execute('SELECT employee_key FROM app_org_employee WHERE manager_key=?', (owner,))}
    return {r[0] for r in conn.execute('''WITH RECURSIVE team(k) AS (
        SELECT employee_key FROM app_org_employee WHERE manager_key=?
        UNION SELECT e.employee_key FROM app_org_employee e JOIN team t ON e.manager_key=t.k)
        SELECT k FROM team WHERE k != ?''', (owner, owner))}


def normalize_filters(raw):
    period = raw.get('period', 'current')
    scope = raw.get('scope', 'direct')
    if period not in {'current', 'week', 'range'} or scope not in {'direct', 'all'}:
        raise ValueError('Некорректный период или охват')
    result = dict(period=period, scope=scope, employees=list(dict.fromkeys(raw.get('employees', []))),
                  surname=str(raw.get('surname', ''))[:250], text=str(raw.get('text', ''))[:500])
    if not isinstance(raw.get('employees', []), list) or len(result['employees']) > 10000:
        raise ValueError('Некорректный выбор сотрудников')
    if period != 'current':
        start = date.fromisoformat(raw.get('date_from', ''))
        end = date.fromisoformat(raw.get('date_to', '')) if period == 'range' else start + timedelta(days=6)
        if period == 'week':
            start -= timedelta(days=start.weekday())
            end = start + timedelta(days=6)
        if end < start:
            raise ValueError('Начало периода должно быть не позже окончания')
        result.update(date_from=start.isoformat(), date_to=end.isoformat())
    return result


def filter_dates(filters, today=None):
    if filters['period'] == 'current':
        today = today or date.today()
        start = today - timedelta(days=today.weekday())
        return start.isoformat(), (start + timedelta(days=6)).isoformat()
    return filters['date_from'], filters['date_to']
